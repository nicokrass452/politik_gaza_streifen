"""German AKUF presentation. Run: pip install python-pptx Pillow requests PyMuPDF
Then: python build_presentation.py
Sources were accessed on 14 October 2026; dates on images are historical.
"""
from pathlib import Path
from io import BytesIO
import json, math, re, argparse
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
OUT = ROOT / 'output'
BUILD = ROOT / 'build'
for p in (ASSETS, OUT, BUILD): p.mkdir(exist_ok=True)
ACCESS = '14.10.2026'
BG, SURFACE, BORDER, TEXT, MUTED, ACCENT = 'F8FAFC','FFFFFF','E2E8F0','0F172A','475569','2563EB'
FONT = 'DejaVu Sans'
SW, SH = 13.333, 7.5
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
prs.core_properties.title = 'Konfliktanalyse: Der Krieg im Gazastreifen (AKUF-Schema)'
prs.core_properties.subject = 'Politik, Klasse 10: Konfliktforschung, UN-Sicherheitsrat und Friedensmodelle'
prs.core_properties.author = '[Name 1], [Name 2], [Name 3]'
prs.core_properties.language = 'de-DE'

SOURCES = {
'akuf': ('AKUF: Kriegsdefinition und Kriegstypologie', 'https://www.wiso.uni-hamburg.de/fachbereich-sowi/ueber-den-fachbereich/fachgebiete/akuf/kriegsdefinition.html'),
'bpb': ('bpb: Der Gazastreifen im Nahostkonflikt (15.03.2024)', 'https://www.bpb.de/shop/zeitschriften/apuz/naher-osten-2024/546377/der-gazastreifen-im-nahostkonflikt/'),
'history': ('bpb: Geschichte des Nahostkonflikts', 'https://www.bpb.de/themen/naher-mittlerer-osten/israel/45042/geschichte-des-nahostkonflikts/'),
'peace': ('bpb: Frieden, Galtung und Senghaas', 'https://www.bpb.de/lernen/politische-bildung/193093/frieden/?p=0'),
'galtung': ('bpb: Überlegungen zur Entspannung in einer pluralistischen Welt', 'https://www.bpb.de/shop/zeitschriften/apuz/archiv/531674/ueberlegungen-zur-entspannung-in-einer-pluralistischen-welt/'),
'swp': ('SWP: Krieg in Nahost', 'https://www.swp-berlin.org/fokusthema/krieg-in-nahost'),
'qatar': ('SWP: Katars Außenpolitik (2022/S 12)', 'https://www.swp-berlin.org/10.18449/2022S12/'),
'map': ('OCHA: Gaza Strip access and movement (Juli 2024)', 'https://www.ochaopt.org/content/gaza-strip-access-and-movement-july-2024'),
'unosat': ('UNOSAT: Gaza Strip Comprehensive Damage Assessment (08.07.2025)', 'https://unosat.org/products/4165'),
'unosatpdf': ('UNOSAT: Originalkarte, veröffentlicht 06.08.2025, Version 1', 'https://unosat.org/static/unosat_filesystem/4165/OCHA-OPT-027_UNOSAT_A3_Gaza_Strip_OPT_CDA_08072025.pdf'),
'unsc': ('UN-Sicherheitsrat: Resolutionen', 'https://main.un.org/securitycouncil/en/content/resolutions-0'),
'r2728': ('UN-Sicherheitsrat: Resolution 2728 (25.03.2024)', 'https://www.un.org/unispal/document/sc-adopted-resolution-2728-2024-demanding-for-an-immediate-ceasefire-during-the-month-of-ramadan-25mar24/'),
'r2735': ('UN-Sicherheitsrat: Resolution 2735 (10.06.2024)', 'https://press.un.org/en/2024/sc15723.doc.htm'),
'r2803': ('UN-Sicherheitsrat: Resolution 2803 (17.11.2025)', 'https://digitallibrary.un.org/record/4093207'),
'r2803ex': ('Vereinte Nationen: Erläuterung der Resolution 2803', 'https://palestine.un.org/en/305520-un-security-council-authorizes-temporary-international-force-gaza'),
'veto1': ('UN: US-Veto vom 18.10.2023', 'https://www.un.org/unispal/document/security-council-fails-to-adopt-resolution-calling-for-humanitarian-pauses-in-israel-gaza-crisis-on-account-of-veto-by-united-states/'),
'veto2': ('UN: Russische und chinesische Vetos vom 25.10.2023', 'https://www.un.org/unispal/document/security-council-fails-to-adopt-either-of-two-draft-resolutions-addressing-conflict-and-humanitarian-crisis-in-gaza-press-release/'),
'charter': ('UN-Charta: vollständiger Text, Artikel 27 und 51', 'https://www.un.org/en/about-us/un-charter/full-text'),
'icj': ('IGH: Südafrika gegen Israel, Verfahren Nr. 192', 'https://www.icj-cij.org/case/192'),
'occupation': ('IGH: Zusammenfassung des Gutachtens vom 19.07.2024', 'https://www.icj-cij.org/node/204176'),
'icc': ('IStGH: Haftbefehle gegen Netanyahu und Gallant (21.11.2024)', 'https://www.icc-cpi.int/news/situation-state-palestine-icc-pre-trial-chamber-i-rejects-state-israels-challenges?embed=true'),
'icc2': ('IStGH: Situation in Palästina, einschließlich Verfahren gegen Deif', 'https://bridge.www.icc-cpi.int/palestine'),
'icrc': ('IKRK: Erklärung zu Gaza und Israel (09.03.2024)', 'https://www.icrc.org/en/document/statement-gaza-and-israel-president-icrc'),
'icrc2': ('IKRK: Fragen und Antworten zu Israel und besetzten Gebieten', 'https://www.icrc.org/en/document/FAQ-icrc-work-Israel-and-occupied-territories'),
'us': ('Congressional Research Service: Abstimmungen zu Israel–Hamas', 'https://www.congress.gov/crs-products/product/pdf/R/R48289'),
'iran': ('Congressional Research Service: Iran’s Foreign and Defense Policies', 'https://www.congress.gov/crs_external_products/R/PDF/R44017/R44017.66.pdf'),
'ocha': ('OCHA: Aufgaben der humanitären Koordination', 'https://www.ochaopt.org/about-us/our-work'),
'unrwa': ('UNRWA: Aufgaben und Hilfsleistungen', 'https://www.unrwa.org/what-we-do'),
'current': ('Reuters: Fortdauernde Gewalt und stockende Umsetzung (12.10.2026)', 'https://www.reuters.com/world/middle-east/israeli-fire-kills-two-people-gaza-medics-say-2026-10-12/'),
'hostages': ('Bericht des Board of Peace, S/2026/418 (22.05.2026)', 'https://www.un.org/unispal/document/implementation-of-united-nations-security-council-resolution-2803-2025-report-of-the-board-of-peace-through-the-office-of-the-high-representative-for-gaza-s-2026-418/'),
}

def ensure_assets():
    import requests, fitz
    jobs = [('ocha_gaza_2024.pdf','https://www.ochaopt.org/sites/default/files/Gaza_A0_2024_july-2024.pdf'),
            ('unosat_2025-07-08.pdf', SOURCES['unosatpdf'][1])]
    for fn, url in jobs:
        p=ASSETS/fn
        if not p.exists():
            response=requests.get(url,timeout=60);response.raise_for_status();p.write_bytes(response.content)
    p=ASSETS/'ocha_map_excerpt.png'
    if not p.exists():
        with fitz.open(ASSETS/'ocha_gaza_2024.pdf') as d:
            d[0].get_pixmap(matrix=fitz.Matrix(1.1,1.1),clip=fitz.Rect(0,210,2376,2800)).save(p)
    p=ASSETS/'unosat_2025-07-08.png'
    if not p.exists():
        with fitz.open(ASSETS/'unosat_2025-07-08.pdf') as d:
            d[0].get_pixmap(matrix=fitz.Matrix(2,2)).save(p)


def color(h): return RGBColor.from_string(h)

def rect(slide,x,y,w,h,fill=SURFACE,line=BORDER):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid();s.fill.fore_color.rgb=color(fill)
    s.text_frame.word_wrap=True
    if line: s.line.color.rgb=color(line);s.line.width=Pt(.7)
    else: s.line.fill.background()
    return s


def txt(slide,text,x,y,w,h,size=18,bold=False,fill=TEXT,align=PP_ALIGN.LEFT):
    s=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=s.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(fill)
        p.alignment=align;p.space_after=Pt(5);p.line_spacing=1.12
        for run in p.runs:
            run._r.get_or_add_rPr().set('lang','de-DE')
    return s


def bullets(slide,items,x,y,w,size=18,gap=.78):
    for i,t in enumerate(items):
        txt(slide,'•',x,y+i*gap,.22,.5,size,False,ACCENT)
        txt(slide,t,x+.32,y+i*gap,w-.32,gap-.07,size,False,TEXT)


def base(title,n,speaker,source=''):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color(BG)
    if title: txt(s,title,.8,.62,11.73,.74,30,True)
    rect(s,.8,7.02,11.73,.013,BORDER,None)
    txt(s,f'[Name {speaker}]  /  Folie {n:02d}',.8,7.10,3.1,.22,10,False,MUTED)
    if source: txt(s,source,4.0,7.10,8.53,.22,9,False,MUTED,PP_ALIGN.RIGHT)
    return s

IMAGE_CHECKS=[]
def add_fitted_image(slide,img_path,box_left,box_top,max_w,max_h):
    with Image.open(img_path) as im:
        im=ImageOps.exif_transpose(im)
        w_px,h_px=im.size
        data=BytesIO();im.save(data,format='PNG');data.seek(0)
    if min(w_px,h_px,max_w,max_h)<=0: raise ValueError('Positive dimensions required')
    aspect=w_px/h_px;box_aspect=max_w/max_h
    if aspect>box_aspect: final_w=max_w;final_h=max_w/aspect
    else: final_h=max_h;final_w=max_h*aspect
    left=box_left+(max_w-final_w)/2;top=box_top+(max_h-final_h)/2
    shape=slide.shapes.add_picture(data,round(left),round(top),width=round(final_w),height=round(final_h))
    assert shape.left>=box_left-1 and shape.top>=box_top-1
    assert shape.left+shape.width<=box_left+max_w+1
    assert shape.top+shape.height<=box_top+max_h+1
    assert abs((shape.width/shape.height)/aspect-1)<1e-5
    IMAGE_CHECKS.append(dict(file=Path(img_path).name,pixels=[w_px,h_px],aspect=aspect,box=[box_left,box_top,max_w,max_h],placed=[shape.left,shape.top,shape.width,shape.height]))
    return shape

NOTES=[]
def note(slide,speaker,seconds,text,keys):
    citations='\n'.join(f'{SOURCES[k][0]}\n{SOURCES[k][1]} [Zugriff am: {ACCESS}]' for k in keys)
    full=f'Sprecher: [Name {speaker}]\nRichtzeit: {seconds} Sekunden\n\nSPRECHTEXT\n{text.strip()}\n\nQUELLEN UND EINORDNUNG\n{citations}'
    slide.notes_slide.notes_text_frame.text=full
    NOTES.append(dict(slide=len(prs.slides),speaker=speaker,seconds=seconds,text=text.strip(),sources=keys))


def card(slide,x,y,w,h,kicker,title,body):
    rect(slide,x,y,w,h)
    txt(slide,kicker,x+.25,y+.23,w-.5,.33,17,True,ACCENT)
    txt(slide,title,x+.25,y+.73,w-.5,.75,18,True)
    txt(slide,body,x+.25,y+1.61,w-.5,h-1.8,15,False,MUTED)

ensure_assets()
# 1
s=base('',1,1,'Projekt Konfliktforschung / Recherche: 14.10.2026')
txt(s,'POLITIK / KLASSE 10',.8,1.12,9,.35,12,True,ACCENT)
txt(s,'Konfliktanalyse:\nDer Krieg im\nGazastreifen',.8,1.80,11.0,2.65,44,True)
txt(s,'AKUF-Schema\nUrsachen, UN-Sicherheitsrat und Wege zum Frieden',.8,4.75,11.2,.90,20,False,MUTED)
txt(s,'[Name 1]    [Name 2]    [Name 3]',.8,6.15,11,.4,17,True)
note(s,1,45,'''Unsere Präsentation untersucht den Krieg im Gazastreifen anhand des AKUF-Schemas. AKUF steht für Arbeitsgemeinschaft Kriegsursachenforschung an der Universität Hamburg. Wir fragen nach Entstehung, Interessen, Mitteln und Möglichkeiten einer Konfliktregelung. Zuerst zeigen wir den Raum und erklären zentrale Begriffe. Danach folgen die Ursachen und eine Satellitenanalyse. Abschließend prüfen wir die Rolle des UN-Sicherheitsrats und formulieren ein eigenes Urteil. Dabei unterscheiden wir zwischen Zivilbevölkerung, bewaffneten Gruppen und Regierungen. Die Verantwortung für konkrete Handlungen muss jeweils einzeln beurteilt werden. Unsere Leitfrage lautet: Welche Bedingungen braucht ein Frieden, der Menschen auf beiden Seiten dauerhaft schützt?''',['akuf','swp'])
# 2
s=base('Raum und Konflikt',2,1,'OCHA: Juli 2024 / bpb: 2024')
add_fitted_image(s,ASSETS/'ocha_map_excerpt.png',Inches(.8),Inches(1.55),Inches(5.4),Inches(4.8))
txt(s,'Quelle: OCHA, Juli 2024, Kartenausschnitt\nRot: damalige Zugangsbeschränkungen',.8,6.42,5.5,.42,10,False,MUTED)
txt(s,'2,1 Mio.',6.75,1.62,5.0,.90,46,True,ACCENT)
txt(s,'Menschen auf 365 km²\n≈ 5.750 Menschen je km²',6.75,2.62,5.0,.82,17,True)
txt(s,'Flächenmittel, OCHA Juli 2024; keine aktuelle Verteilungskarte',6.75,3.49,5.0,.43,10,False,MUTED)
bullets(s,['Israel / IDF gegen Hamas und Palästinensischen Islamischen Dschihad','Konflikt um Sicherheit, Territorium und palästinensische Selbstbestimmung','Nachbarn: Israel und Ägypten; Städte: Gaza, Khan Yunis, Rafah'],6.75,4.16,5.45,16,.76)
note(s,1,70,'''Der Gazastreifen liegt am Mittelmeer. Im Norden und Osten grenzt er an Israel, im Süden an Ägypten. Auf der Karte erkennen wir Gaza-Stadt, Khan Yunis und Rafah. OCHA nennt für Juli 2024 ungefähr 2,1 Millionen Menschen auf 365 Quadratkilometern. Daraus ergibt sich rechnerisch ein Flächenmittel von etwa 5.750 Menschen je Quadratkilometer. Die tatsächliche Dichte unterscheidet sich stark, besonders durch Vertreibung. Die Karte zeigt die Zugangslage von Juli 2024 und keine heutigen sicheren Wege. Zu den zentralen bewaffneten Akteuren gehören die israelischen Streitkräfte, kurz IDF, sowie Hamas und der Palästinensische Islamische Dschihad. Die palästinensische Bevölkerung ist nicht mit diesen Gruppen gleichzusetzen. Hinter den Kämpfen stehen konkurrierende Ansprüche auf Sicherheit, Land und politische Selbstbestimmung. Für den weiteren Vortrag ist deshalb wichtig, zwischen den Zielen der Bevölkerung und denen politischer oder militärischer Führungen zu unterscheiden.''',['map','bpb','swp'])
# 3
s=base('Glossar: fünf Begriffe',3,1,'UN-Charta: Art. 27 und Kapitel VII / IKRK / bpb')
w=(SW-1.6-.4)/2
for x,y,title,body in [
(.8,1.72,'Asymmetrischer Konflikt','Ungleiche militärische Stärke und unterschiedliche Methoden der Kriegsführung.'),
(.8+w+.4,1.72,'Humanitäres Völkerrecht','Regeln im Krieg, etwa die Genfer Konventionen. Sie schützen Zivilisten und begrenzen Gewalt.'),
(.8,4.10,'UN-Sicherheitsrat & Veto','15 Mitglieder, fünf ständige Vetomächte. Art. 27 regelt Abstimmungen; Kapitel VII ermöglicht Zwangsmaßnahmen.'),
(.8+w+.4,4.10,'Zwei-Staaten-Lösung / Blockade','Zwei souveräne Staaten für Israelis und Palästinenser. Blockade: starke Einschränkung des Zugangs und Warenverkehrs.')]:
    rect(s,x,y,w,2.10);txt(s,title,x+.24,y+.23,w-.48,.68,18,True);txt(s,body,x+.24,y+1.02,w-.48,.88,14,False,MUTED)
note(s,1,65,'''Ein asymmetrischer Konflikt bedeutet, dass die Gegner sehr unterschiedliche militärische Möglichkeiten haben. Das sagt für sich genommen noch nichts darüber aus, welche Handlung rechtmäßig ist. Das humanitäre Völkerrecht gilt für alle Konfliktparteien. Es schützt zum Beispiel Menschen, die nicht oder nicht mehr kämpfen. Der UN-Sicherheitsrat trägt die Hauptverantwortung für den Weltfrieden. Bei Sachentscheidungen sind grundsätzlich neun Stimmen nötig, und keine der fünf ständigen Mächte darf mit Nein stimmen. Eine Enthaltung ist kein Veto. Das Vetorecht steht in Artikel 27, also nicht in Kapitel VII. Kapitel VII betrifft Maßnahmen bei Friedensbedrohungen. Die Zwei-Staaten-Lösung meint Israel und einen souveränen palästinensischen Staat. Die Blockade beschränkt dagegen die Bewegungsfreiheit und den Warenverkehr. Sicherheitsbegründungen und die Folgen für die Zivilbevölkerung müssen wir dabei getrennt untersuchen.''',['charter','icrc','bpb'])
# 4
s=base('Genese: fünf historische Einschnitte',4,1,'bpb: Geschichte des Nahostkonflikts / Gazastreifen, 2024')
txt(s,'Der Krieg seit 2023 hat eine längere Vorgeschichte',.8,1.53,11.73,.5,20,False,MUTED)
line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(.92),Inches(2.80),Inches(12.40),Inches(2.80));line.line.color.rgb=color(BORDER);line.line.width=Pt(2)
items=[('1948','Israel / Nakba','Staatsgründung, Krieg, Flucht und Vertreibung'),('1967','Besetzung','Israel besetzt Gaza und weitere Gebiete'),('2005','Abzug','Israel räumt Siedlungen und zieht Truppen ab'),('2007','Hamas / Blockade','Machtübernahme; verschärfte Abriegelung'),('07.10.2023','Terrorangriff','Hamas-geführter Angriff; Beginn der neuen Kriegsphase')]
for i,(date,title,body) in enumerate(items):
    x=.8+i*2.41
    txt(s,date,x,2.22,2.18,.43,19,True,ACCENT)
    rect(s,x,2.74,.12,.12,ACCENT,None)
    txt(s,title,x,3.15,2.15,.77,18,True)
    txt(s,body,x,4.03,2.15,1.32,15,False,MUTED)
rect(s,.8,5.85,11.73,.64)
txt(s,'Geschichte erklärt Konfliktlinien. Sie rechtfertigt keine Angriffe auf Zivilisten.',1.02,6.02,11.3,.31,15,True)
note(s,1,60,'''1948 wurde Israel gegründet. Im anschließenden Krieg flohen Hunderttausende Palästinenser oder wurden vertrieben. Dafür steht der Begriff Nakba, auf Deutsch Katastrophe. Gaza kam zunächst unter ägyptische Verwaltung. Im Sechstagekrieg 1967 besetzte Israel den Gazastreifen und weitere Gebiete. 2005 zog Israel seine Truppen aus Gaza ab und räumte die Siedlungen. Der Abzug bedeutete jedoch keine palästinensische Souveränität. Nach innerpalästinensischen Kämpfen übernahm Hamas 2007 die Kontrolle. Israel verschärfte die Blockade, auch Ägypten beschränkte den Zugang. Am 7. Oktober 2023 griffen Hamas und weitere bewaffnete Gruppen Israel an, ermordeten Zivilisten und verschleppten Menschen. Israel begann anschließend einen großangelegten Krieg in Gaza. Diese Geschichte erklärt Hintergründe und Ängste, aber sie entschuldigt keine Gewalt gegen Unbeteiligte. Nun übernimmt Name 2 mit den Interessen und Mitteln der Akteure.''',['history','bpb','occupation','swp'])
# 5
s=base('Interessen der Konfliktparteien',5,2,'Eigene Gegenüberstellung auf Basis von bpb und SWP')
w=(SW-1.6-.8)/3
card(s,.8,1.75,w,4.35,'01','Israelische Regierung','Schutz und Sicherheit Israels\nHamas militärisch entmachten\nPolitische Macht erhalten')
card(s,.8+w+.4,1.75,w,4.35,'02','Palästinensische Gesellschaft / PA','Selbstbestimmung und Staatlichkeit\nSchutz, Bewegungsfreiheit, Wiederaufbau\nPolitische Teilhabe')
card(s,.8+2*(w+.4),1.75,w,4.35,'03','Hamas-Führung','Herrschaft und Einfluss sichern\nBewaffneter Kampf gegen Israel\nIslamistischer Ordnungsanspruch')
txt(s,'Streitpunkte: Grenzen, Territorium und Sicherheit; Ziele innerhalb beider Gesellschaften sind umstritten.',.8,6.37,11.73,.45,14,False,MUTED)
note(s,2,60,'''Die Interessen lassen sich nicht einfach auf zwei einheitliche Lager reduzieren. Die israelische Regierung begründet ihr Vorgehen mit Sicherheit und der Ausschaltung der militärischen Bedrohung durch Hamas. Innerhalb Israels gibt es jedoch Streit über Kriegsführung, territoriale Kontrolle und politische Lösungen. Viele Palästinenser verlangen einen eigenen Staat, Bewegungsfreiheit und ein Leben ohne Besatzung und Gewalt. Die Palästinensische Autonomiebehörde, kurz PA, ist eine andere politische Organisation als Hamas. Hamas verbindet den bewaffneten Kampf gegen Israel mit dem Anspruch, ihre eigene Herrschaft und eine islamistische Ordnung zu erhalten. Auch palästinensische Stimmen kritisieren diese Herrschaft. Bei politischen Führungen können außerdem Machterhalt und ideologische Ansprüche Kompromisse erschweren. Unsere Gegenüberstellung ist eine analytische Zusammenfassung, keine Behauptung über die Meinung jedes einzelnen Menschen. Sicherheit für Israelis und Selbstbestimmung für Palästinenser müssen in einer tragfähigen Lösung gemeinsam berücksichtigt werden.''',['bpb','swp','qatar'])
# 6
s=base('Rechtliche Situation und Machtmittel',6,2,'UN-Charta / IGH / IStGH / IKRK / CRS / SWP')
txt(s,'Rechtliche Situation',.8,1.65,5.55,.4,21,True,ACCENT)
bullets(s,['Art. 51: Selbstverteidigung unter rechtlichen Voraussetzungen','Kriegsvölkerrecht bindet alle Parteien und schützt Zivilisten','IGH: Staatenverantwortung; Eilmaßnahmen im Verfahren Südafrika gegen Israel','IStGH: individuelle Strafverantwortung; Haftbefehle 2024'],.8,2.35,5.55,17,.89)
txt(s,'Mittel und Ressourcen',6.88,1.65,5.55,.4,21,True,ACCENT)
bullets(s,['Israel: überlegene Streitkräfte; insbesondere militärische US-Unterstützung','Hamas / PIJ: Raketen, Tunnel, asymmetrische Kriegsführung','Iran: Unterstützung bewaffneter Gruppen der „Achse des Widerstands“','Katar: politische Kontakte, Finanzhilfen und Vermittlung'],6.88,2.35,5.55,17,.89)
txt(s,'Eilmaßnahmen und Haftbefehle sind keine rechtskräftigen Schuldsprüche.',.8,6.32,11.73,.47,16,True)
note(s,2,70,'''Wir müssen zwei rechtliche Fragen auseinanderhalten: Unter welchen Voraussetzungen darf ein Staat Gewalt einsetzen, und welche Regeln gelten während eines Krieges? Israel beruft sich auf Selbstverteidigung nach Artikel 51 der UN-Charta. Dessen Anwendung im konkreten Besatzungskontext ist juristisch umstritten. Unabhängig davon gelten die Regeln zum Schutz der Zivilbevölkerung, zur Unterscheidung militärischer Ziele und zur Verhältnismäßigkeit. Geiselnahmen sind verboten. Der Internationale Gerichtshof behandelt die Verantwortung von Staaten. Im Verfahren Südafrika gegen Israel ordnete er 2024 vorläufige Schutzmaßnahmen an. Daraus folgt noch kein abschließendes Urteil über den Völkermordvorwurf. Der Internationale Strafgerichtshof befasst sich dagegen mit einzelnen Personen. Er erließ 2024 Haftbefehle gegen Netanyahu, Gallant und den Hamas-Kommandeur Deif. Haftbefehle sind keine Verurteilungen. Militärisch ist Israel weit überlegen und erhält insbesondere US-Unterstützung. Hamas und PIJ nutzen unter anderem Raketen und Tunnel. Iran unterstützt bewaffnete Partner. Katar unterhält Kontakte zu Hamas, finanziert Hilfen und vermittelt. Diese Rollen dürfen nicht gleichgesetzt werden.''',['charter','icrc','icrc2','occupation','icj','icc','icc2','us','iran','qatar'])
# 7
s=base('Eskalation und sichtbare Zerstörung',7,2,'UNOSAT: Bildstand 08.07.2025 / SWP / Reuters, 12.10.2026')
txt(s,'Eskalationsstufen',.8,1.65,5.7,.40,21,True,ACCENT)
bullets(s,['07.10.2023: Terrorangriff und Geiselnahmen','Ab Oktober 2023: Bombardierung und Bodenoffensive','Regionale Ausweitung: Hisbollah, Huthi und Iran','Seit Oktober 2025: Waffenruhe mit fortdauernder Gewalt'],.8,2.23,5.5,16,.65)
txt(s,'78 %',.8,5.08,2.5,.80,42,True,ACCENT)
txt(s,'Bauwerke betroffen*\n102.067 davon zerstört',3.07,5.18,3.39,.7,14,True)
txt(s,'* Einschließlich möglicher Schäden; vorläufige Auswertung.\nNicht vor Ort verifiziert; keine Aussage zu Opferzahlen oder Tätern.',.8,6.19,5.53,.55,10,False,MUTED)
add_fitted_image(s,ASSETS/'unosat_2025-07-08.png',Inches(6.91),Inches(1.70),Inches(5.4),Inches(4.5))
txt(s,'Quelle: UNOSAT, 08.07.2025; Satellitendaten:\nWorldView-2/-3 (Maxar), Pléiades-NEO (CNES)\nRot zeigt die Dichte erfasster Schäden.',6.91,6.30,5.4,.62,10,False,MUTED)
note(s,2,60,'''Nach dem Terrorangriff vom 7. Oktober 2023 folgten israelische Luftangriffe und die Bodenoffensive. Zugleich weitete sich der Konflikt durch Hisbollah, Huthi und direkte Auseinandersetzungen mit Iran regional aus. Die Waffenruhe vom Oktober 2025 führte nicht zu einem dauerhaften Ende der Gewalt. Für unsere Bildanalyse betrachten wir bewusst einen früheren, klar datierten Stand: den 8. Juli 2025. UNOSAT erfasste damals 192.812 betroffene Bauwerke. Dazu gehören zerstörte, schwer oder mäßig beschädigte sowie möglicherweise beschädigte Strukturen. Insgesamt entspricht das rund 78 Prozent. 102.067 Bauwerke wurden als zerstört eingestuft. Die roten Flächen zeigen die räumliche Dichte der Schäden. Das hilft, den Umfang der Zerstörung und den Wiederaufbaubedarf zu verstehen. Die Auswertung war vorläufig und noch nicht vor Ort überprüft. Aus dem Satellitenbild allein lassen sich weder Opferzahlen noch die Verantwortlichkeit für einzelne Angriffe oder sichere Fluchtwege ableiten.''',['swp','current','unosat','unosatpdf'])
# 8
s=base('AKUF-Typologie: eine begründete Einordnung',8,2,'Eigene Analyse nach AKUF; keine offizielle AKUF-Fallcodierung')
txt(s,'Die fünf Typen',.8,1.65,5.1,.4,21,True,ACCENT)
for i,(a,b) in enumerate([('A','Antiregimekrieg'),('B','Autonomie- / Sezessionskrieg'),('C','Zwischenstaatlicher Krieg'),('D','Dekolonisationskrieg'),('E','Sonstige Kriege')]):
    txt(s,a,.8,2.32+i*.61,.45,.4,21,True,ACCENT);txt(s,b,1.42,2.34+i*.61,4.7,.39,17)
rect(s,6.68,1.67,5.85,4.45)
txt(s,'Mischtyp als Arbeitsthese',6.94,1.93,5.30,.75,22,True)
txt(s,'A: Kampf um die Herrschaft der Hamas\nB: Dimension politischer Selbstbestimmung\nC: besonders bei direktem Israel–Iran-Krieg',6.94,3.0,5.12,1.6,17,False,MUTED)
txt(s,'Besatzung ist kein eigener AKUF-Typ.\nDie Einordnung hängt von Zeitraum und Konfliktebene ab.',6.94,4.98,5.13,.87,14,True)
txt(s,'Fremdbeteiligung: direkte Kämpfe fremder staatlicher Truppen; Waffen oder Geld allein reichen nicht.',.8,6.37,11.73,.46,14,False,MUTED)
note(s,2,50,'''AKUF unterscheidet fünf Typen. A bezeichnet Antiregimekriege, B Autonomie- und Sezessionskriege, C zwischenstaatliche Kriege, D Dekolonisationskriege und E sonstige Kriege. Für unsere Analyse wählen wir eine vorsichtige Mischtyp-These. A erfasst den Kampf um das Ende oder den Erhalt der Hamas-Herrschaft. B passt teilweise zum Streben nach Selbstbestimmung, aber Gaza ist kein gewöhnlicher Sezessionsfall. Eine klare C-Dimension entsteht vor allem bei direkten Kämpfen zwischen Israel und Iran. Die De-facto-Regierung der Hamas erschwert eine einfache Trennung zwischen staatlichen und nichtstaatlichen Akteuren. Besatzung beschreibt zudem einen rechtlichen Kontext und keinen eigenen Buchstabentyp. Wir behaupten deshalb keine offizielle AKUF-Codierung für Gaza. Fremdbeteiligung zählt bei AKUF nur bei unmittelbarer Teilnahme fremder staatlicher Streitkräfte. Waffenlieferungen oder Geld genügen nicht. Nun untersucht Name 3 die internationalen Handlungsmöglichkeiten.''',['akuf','occupation','swp'])
# 9
s=base('UN-Sicherheitsrat: Blockaden und Beschlüsse',9,3,'UN-Sicherheitsrat: 2728 / 2735 / 2803; Abstimmungen 2023–2025')
txt(s,'18.10.2023: US-Veto  /  25.10.2023: Vetos Russlands und Chinas',.8,1.50,11.73,.49,17,True)
w=(SW-1.6-.8)/3
for i,(num,date,title,body) in enumerate([
('2728','25.03.2024','Ramadan-Waffenruhe','Sofortige Waffenruhe und bedingungslose Freilassung aller Geiseln'),
('2735','10.06.2024','Drei-Phasen-Plan','Waffenruhe und Austausch; dauerhaftes Ende der Kämpfe; Wiederaufbau'),
('2803','17.11.2025','Übergangsordnung','Umfassender Plan; Board of Peace; internationale Stabilisierungstruppe (ISF)')]):
    x=.8+i*(w+.4);rect(s,x,2.30,w,3.24)
    txt(s,num,x+.24,2.54,w-.48,.55,28,True,ACCENT)
    txt(s,date,x+.24,3.14,w-.48,.3,12,False,MUTED)
    txt(s,title,x+.24,3.72,w-.48,.7,18,True)
    txt(s,body,x+.24,4.50,w-.48,.84,14,False,MUTED)
txt(s,'Vermittler: USA, Ägypten, Katar',.8,5.95,11.7,.36,17,True)
txt(s,'UNRWA: Versorgung und Bildung  /  OCHA: Koordination humanitärer Hilfe',.8,6.46,11.73,.36,14,False,MUTED)
note(s,3,80,'''Der Sicherheitsrat kann verbindliche Entscheidungen treffen, bleibt aber von den Interessen seiner Mitglieder abhängig. Am 18. Oktober 2023 blockierten die USA einen brasilianischen Entwurf für humanitäre Pausen. Eine Woche später legten Russland und China Vetos gegen einen amerikanischen Entwurf ein. Die Gründe unterschieden sich, unter anderem bei Selbstverteidigung und der Forderung nach einem Waffenstillstand. Am 25. März 2024 verlangte Resolution 2728 eine sofortige Waffenruhe für Ramadan und die bedingungslose Freilassung aller Geiseln. Die USA enthielten sich. Resolution 2735 vom Juni unterstützte einen Plan mit drei Phasen: zunächst Waffenruhe und Austausch, dann ein dauerhaftes Ende der Kämpfe und schließlich Wiederaufbau. Resolution 2803 vom November 2025 unterstützte den umfassenden Plan und begrüßte den Board of Peace als Übergangsverwaltung. Sie autorisierte eine zeitweilige internationale Stabilisierungstruppe. Russland und China enthielten sich diesmal. Ein beschlossenes Mandat bedeutet jedoch nicht automatisch erfolgreiche Umsetzung. Als Vermittler wirken USA, Ägypten und Katar. UNRWA stellt grundlegende Hilfsleistungen bereit, während OCHA humanitäre Hilfe koordiniert. Diese Organisationen sind keine Streitkräfte des Sicherheitsrats.''',['veto1','veto2','r2728','r2735','r2803','r2803ex','qatar','unrwa','ocha'])
# 10
s=base('Kompromisse und Wege zu einer Lösung',10,3,'Eigene Ableitung aus UN-Resolutionen, IKRK und SWP')
for x,y,title,body in [
(.8,1.73,'Anreize für Verhandlungen','Geiselabkommen als bisheriger Hebel; humanitäre Not, militärische Kosten und internationaler Druck'),
(6.867,1.73,'Überprüfbare Vereinbarungen','Verlässliche Waffenruhe, Schutz von Zivilisten, Hilfe und überprüfte Sicherheitszusagen'),
(.8,4.11,'Politische Perspektive','Legitime palästinensische Vertretung, Wiederaufbau und Sicherheit für Israel; Perspektive zweier Staaten'),
(6.867,4.11,'Zivilgesellschaft','Opferverbände, Friedensinitiativen und lokale Hilfsorganisationen einbeziehen; Gewalt und Entmenschlichung zurückweisen')]:
    rect(s,x,y,5.666,2.10);txt(s,title,x+.24,y+.23,5.18,.7,18,True);txt(s,body,x+.24,y+1.04,5.18,.87,14,False,MUTED)
note(s,3,60,'''Verhandlungen werden wahrscheinlicher, wenn beide Seiten Vorteile erwarten und die Kosten einer Fortsetzung steigen. Geisel- und Gefangenenaustausch waren dafür ein wichtiger Hebel. Das ist hier historisch gemeint: Der Bericht des Board of Peace nennt die Rückkehr der letzten lebenden Geisel im Oktober 2025 und der letzten verstorbenen Geisel im Januar 2026. Heute stehen die Umsetzung der Vereinbarungen und die politische Zukunft im Zentrum. Humanitäre Hilfe muss die Menschen erreichen und darf nicht von politischen Gegenleistungen abhängig gemacht werden. Für belastbare Abkommen braucht es überprüfbare Sicherheitszusagen und Regeln für Verstöße. Langfristig gehören Wiederaufbau, legitime palästinensische Mitbestimmung und Sicherheit für Israel zusammen. Friedensinitiativen und lokale Organisationen können Vertrauen schaffen, ersetzen aber keine politischen Entscheidungen. Unsere Schlussfolgerung lautet deshalb: Kurzfristige Absprachen müssen mit einer glaubwürdigen Perspektive verbunden sein. Sonst bleiben die Ursachen des Konflikts bestehen.''',['hostages','r2735','r2803ex','icrc','swp','peace'])
# 11
s=base('Unser Urteil: mehr als eine Waffenruhe',11,3,'Eigene Bewertung nach Galtung und Senghaas / Lagebericht: Reuters, 12.10.2026')
txt(s,'Galtung: Maßstab Frieden',.8,1.57,5.55,.42,21,True,ACCENT)
txt(s,'Negativer Frieden',.8,2.23,5.55,.38,19,True)
txt(s,'Direkte Gewalt endet. Eine verletzte Waffenruhe erfüllt dieses Ziel nur unvollständig.',.8,2.80,5.48,.84,17,False,MUTED)
txt(s,'Positiver Frieden',.8,3.90,5.55,.39,19,True)
txt(s,'Sicherheit, Rechte und Selbstbestimmung; Abbau struktureller Gewalt.',.8,4.46,5.48,.84,17,False,MUTED)
txt(s,'Senghaas: offene Bedingungen',6.88,1.57,5.55,.42,21,True,ACCENT)
for i,t in enumerate(['Gewaltmonopol mit rechtsstaatlicher Kontrolle','Rechtsstaatlichkeit und politische Teilhabe','Soziale Gerechtigkeit und Wiederaufbau','Konfliktkultur, Interdependenz und Affektkontrolle']):
    txt(s,t,6.88,2.23+i*.74,5.55,.62,16,False,MUTED)
rect(s,.8,5.87,11.73,.88)
txt(s,'Unser Urteil: Resolution 2803 bietet einen Übergangsrahmen.\nDauerhafter Frieden braucht gleiche Rechte und verlässlichen Schutz für alle.',1.04,6.06,11.24,.57,16,True)
note(s,3,85,'''Unser Urteil richtet sich nach offengelegten Kriterien. Bei Johan Galtung bedeutet negativer Frieden zunächst die Abwesenheit direkter Gewalt. Eine Waffenruhe kann ein Schritt dorthin sein. Wenn trotzdem Menschen getötet werden, ist bereits dieser Mindeststandard unvollständig erfüllt. Reuters berichtete noch am 12. Oktober 2026 über tödliche Angriffe und eine stockende Umsetzung des Abkommens. Positiver Frieden verlangt zusätzlich, strukturelle Gewalt abzubauen und Rechte, Sicherheit und Entwicklung zu ermöglichen. Senghaas nennt sechs miteinander verbundene Bedingungen. Ein Gewaltmonopol braucht rechtsstaatliche Kontrolle. Demokratische Beteiligung gibt Menschen Einfluss auf Entscheidungen. Soziale Gerechtigkeit setzt Zugang zu grundlegender Versorgung voraus. Konstruktive Konfliktkultur sowie Interdependenz und Affektkontrolle helfen, Streit ohne Gewalt zu bearbeiten. Für Gaza sehen wir besonders große Defizite bei legitimer Gewaltkontrolle, Rechtsschutz, Versorgung und politischer Beteiligung. Unser begründetes Urteil lautet: Resolution 2803 ist ein möglicher Übergangsrahmen, aber kein Nachweis eines erreichten Friedens. Entscheidend sind ihre tatsächlichen Folgen für die Menschen. Eine tragfähige Lösung muss Israelis vor Angriffen schützen und Palästinensern Selbstbestimmung und gleiche Rechte ermöglichen. Die Aussicht auf Stabilisierung spricht für einen Übergangsrahmen. Fehlende demokratische Legitimation und unzureichender Schutz wären dagegen gewichtige Einwände. Menschenrechte gelten dabei unabhängig von Herkunft oder Religion.''',['galtung','peace','current','r2803ex','swp'])
# 12
s=base('Quellen und Nachweise',12,3,'[Zugriff am: 14.10.2026] / Ergänzende Belege in den Sprechernotizen')
txt(s,'Alle folgenden Quellen: [Zugriff am: 14.10.2026]',.8,1.47,11.73,.4,15,True,ACCENT)
# Complete URLs are split at path boundaries for legibility; hyperlinks retain the original URLs.
entries=[
('UN-Sicherheitsrat: Resolutionen','https://main.un.org/securitycouncil/en/content/resolutions-0'),
('2728 (2024): Wortlaut','https://docs.un.org/S/RES/2728(2024)'),
('2735 (2024): Wortlaut','https://docs.un.org/S/RES/2735(2024)'),
('2803 (2025): Dokument und Abstimmung','https://digitallibrary.un.org/record/4093207'),
('UNOSAT: Schadensanalyse, 08.07.2025','https://unosat.org/products/4165'),
('AKUF Universität Hamburg: Typologie',SOURCES['akuf'][1]),
('bpb: Der Gazastreifen im Nahostkonflikt',SOURCES['bpb'][1]),
('SWP: Krieg in Nahost',SOURCES['swp'][1]),
('OCHA: Karte, Juli 2024',SOURCES['map'][1]),
('bpb: Frieden und Senghaas',SOURCES['peace'][1]),
]
for i,(label,url) in enumerate(entries):
    col=0 if i<5 else 1;row=i%5;x=.8+6.05*col;y=2.10+row*.88
    txt(s,label,x,y,5.65,.29,12,True)
    # Allow wrapping after URL separators without altering the actual hyperlink.
    display=url.replace('/', '/\u200b').replace('-', '-\u200b')
    shape=txt(s,display,x,y+.33,5.65,.48,9.5,False,MUTED)
    for p in shape.text_frame.paragraphs:
        for r in p.runs: r.hyperlink.address=url
note(s,3,15,'''Hier stehen die zentralen Quellen mit dem tatsächlichen Zugriffsdatum. In den Sprechernotizen finden sich die ergänzenden juristischen Belege und Bildnachweise. Unsere eigene Einordnung und unser Urteil haben wir ausdrücklich gekennzeichnet. Vielen Dank.''',list(SOURCES.keys()))

# Assert required structure and image geometry before writing.
assert len(prs.slides)==12
assert [n['speaker'] for n in NOTES]==[1]*4+[2]*4+[3]*4
assert [sum(n['seconds'] for n in NOTES if n['speaker']==k) for k in (1,2,3)]==[240]*3
for i,slide in enumerate(prs.slides,1):
    assert slide.notes_slide.notes_text_frame.text.strip()
    for shape in slide.shapes:
        assert shape.left>=0 and shape.top>=0,(i,shape.name)
        assert shape.left+shape.width<=prs.slide_width+2,(i,shape.name,'right')
        assert shape.top+shape.height<=prs.slide_height+2,(i,shape.name,'bottom')
        if shape.has_text_frame: assert shape.text_frame.word_wrap is True
parser=argparse.ArgumentParser(description='Erstellt die AKUF-Präsentation mit Sprechernotizen.')
parser.add_argument('--draft', action='store_true', help='Schreibt eine Arbeitsversion für die Layoutprüfung.')
args=parser.parse_args()
candidate=BUILD/'candidate.pptx' if args.draft else OUT/'akuf_gazastreifen_praesentation.pptx'
prs.save(candidate)
(BUILD/'image_checks.json').write_text(json.dumps(IMAGE_CHECKS,indent=2))
(BUILD/'notes_manifest.json').write_text(json.dumps(NOTES,ensure_ascii=False,indent=2))
(ROOT/'quellen.md').write_text('# Quellen\n\nAlle Quellen: [Zugriff am: '+ACCESS+']\n\n'+'\n\n'.join(f'## {name}\n{url}' for name,url in SOURCES.values()))
print('Created:',candidate)
for k in (1,2,3):
    count=sum(len(n['text'].split()) for n in NOTES if n['speaker']==k)
    print('Speaker',k,':',count,'words, target 4 minutes')
print('Images:',len(IMAGE_CHECKS),'all proportional and inside boxes')

"""Baut die Präsentation "Der Krieg im Gazastreifen" (12 Folien, Politik Klasse 10).
Vorher generate_graphics.py und rasterize_svgs.py ausführen, damit die Karten vorliegen.
"""
from pathlib import Path
import json
import argparse

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
BUILD = ROOT / 'build'
OUT = ROOT / 'output'
BUILD.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)

FONT = 'Calibri'
BLUE = RGBColor(0x1F, 0x38, 0x64)      # Überschriften
ACCENT = RGBColor(0x2E, 0x75, 0xB6)    # Linien, Hervorhebungen
TEXT = RGBColor(0x26, 0x26, 0x26)
GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xDE, 0xEB, 0xF7)     # helle Kästen
RULE = RGBColor(0xBF, 0xBF, 0xBF)

SLIDE_W, SLIDE_H = 13.333, 7.5
TOTAL = 12

prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
prs.core_properties.title = 'Der Krieg im Gazastreifen'
prs.core_properties.author = '[Name 1], [Name 2], [Name 3]'
prs.core_properties.subject = 'Politik Klasse 10, Projekt Konfliktforschung (AKUF-Schema)'
prs.core_properties.language = 'de-DE'

NOTES = json.loads((ASSETS / 'speaker_notes.json').read_text(encoding='utf-8'))
IMAGE_CHECKS = []
SVG_SERIAL = 0


# ---------------------------------------------------------------- Hilfsfunktionen

def _style_run(run, size, bold=False, italic=False, color=TEXT):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run._r.get_or_add_rPr().set('lang', 'de-DE')


def textbox(slide, x, y, w, h, text, size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT,
            italic=False, anchor=MSO_ANCHOR.TOP, line_spacing=1.0, space_after=4):
    """Einfacher Textrahmen. Zeilenumbrüche im Text ergeben eigene Absätze."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = line
        _style_run(run, size, bold, italic, color)
    return box


def _add_bullet(paragraph, level=0, numbered=False):
    pPr = paragraph._p.get_or_add_pPr()
    indent = 0.3
    pPr.set('marL', str(int(Inches(indent * (level + 1)))))
    pPr.set('indent', str(int(-Inches(indent))))
    if numbered:
        auto = OxmlElement('a:buAutoNum')
        auto.set('type', 'arabicPeriod')
        pPr.append(auto)
    else:
        bu_font = OxmlElement('a:buFont')
        bu_font.set('typeface', 'Arial')
        bu_char = OxmlElement('a:buChar')
        bu_char.set('char', '•' if level == 0 else '–')
        pPr.append(bu_font)
        pPr.append(bu_char)


def bullets(slide, x, y, w, h, items, size=18, color=TEXT, numbered=False, space_after=8,
            sub_size=None, line_spacing=1.0, space_before=0):
    """Echte PowerPoint-Aufzählung.

    items: Liste aus Strings oder Tupeln. Ein Tupel (fett, normal) ergibt einen fetten
    Anfang. Ein Eintrag der Form ['sub', ...] wird als Unterpunkt eingerückt.
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    first = True
    for item in items:
        level = 0
        if isinstance(item, list) and item and item[0] == 'sub':
            level = 1
            item = item[1]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after if level == 0 else space_after * 0.6)
        if level == 0 and space_before:
            p.space_before = Pt(space_before)
        _add_bullet(p, level, numbered and level == 0)
        this_size = size if level == 0 else (sub_size or size - 2)
        if isinstance(item, tuple):
            lead, rest = item
            r = p.add_run()
            r.text = lead
            _style_run(r, this_size, True, False, color)
            r2 = p.add_run()
            r2.text = rest
            _style_run(r2, this_size, False, False, color)
        else:
            r = p.add_run()
            r.text = item
            _style_run(r, this_size, False, False, color)
    return box


def rect(slide, x, y, w, h, fill=LIGHT, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.text_frame.word_wrap = True
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    _drop_theme_style(shape)
    return shape


def hline(slide, x1, y1, x2, y2, color=RULE, width=0.75):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    _drop_theme_style(conn)
    return conn


def _drop_theme_style(shape):
    """Theme-Stil (inkl. Schatten) entfernen, damit nur die gesetzten Farben gelten."""
    style = shape._element.find(qn('p:style'))
    if style is not None:
        shape._element.remove(style)


def picture(slide, path, x, y, w, h, svg=None):
    """Bild proportional in den Rahmen einpassen; optional das SVG als Vektor-Variante einbetten."""
    global SVG_SERIAL
    with Image.open(path) as im:
        pw, ph = im.size
    scale = min(w / pw, h / ph)
    fw, fh = pw * scale, ph * scale
    left, top = x + (w - fw) / 2, y + (h - fh) / 2
    pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(fw), height=Inches(fh))
    if svg:
        SVG_SERIAL += 1
        part = Part(PackURI(f'/ppt/media/map_{SVG_SERIAL}.svg'), 'image/svg+xml', slide.part.package,
                    blob=Path(svg).read_bytes())
        rid = slide.part.relate_to(part, RT.IMAGE)
        blip = pic._pic.blipFill.blip
        extlst = blip.find(qn('a:extLst'))
        if extlst is None:
            extlst = OxmlElement('a:extLst')
            blip.append(extlst)
        ext = OxmlElement('a:ext')
        ext.set('uri', '{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
        ns = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'
        svg_blip = etree.Element(f'{{{ns}}}svgBlip', nsmap={'asvg': ns})
        svg_blip.set(qn('r:embed'), rid)
        ext.append(svg_blip)
        extlst.append(ext)
    IMAGE_CHECKS.append({'slide': len(prs.slides), 'file': Path(path).name, 'box': [x, y, w, h],
                         'placed': [left, top, fw, fh], 'has_svg': bool(svg)})
    return pic


def map_image(slide, name, x, y, w, h):
    return picture(slide, ASSETS / f'{name}.png', x, y, w, h, ASSETS / f'{name}.svg')


def new_slide(title, number):
    """Leere Folie mit Überschrift, Linie und Fußzeile."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if title:
        textbox(slide, 0.6, 0.35, 12.1, 0.8, title, size=30, bold=True, color=BLUE, anchor=MSO_ANCHOR.BOTTOM)
        hline(slide, 0.65, 1.22, 12.7, 1.22, ACCENT, 1.5)
    textbox(slide, 0.6, 7.0, 6, 0.3, 'Der Krieg im Gazastreifen  |  Politik 10', size=10, color=GREY)
    textbox(slide, 11.7, 7.0, 1.0, 0.3, f'{number} / {TOTAL}', size=10, color=GREY, align=PP_ALIGN.RIGHT)
    return slide


def set_notes(slide, number):
    slide.notes_slide.notes_text_frame.text = NOTES[number - 1]


def caption(slide, x, y, w, text):
    return textbox(slide, x, y, w, 0.3, text, size=10, color=GREY, italic=True)


# ---------------------------------------------------------------- Folien

# 1 Titel
s = new_slide('', 1)
textbox(s, 0.9, 1.6, 11.5, 1.6, 'Der Krieg im Gazastreifen', size=48, bold=True, color=BLUE,
        anchor=MSO_ANCHOR.BOTTOM)
textbox(s, 0.9, 3.25, 11.5, 0.6, 'Eine Konfliktanalyse nach dem AKUF-Schema', size=26, color=GREY)
hline(s, 0.95, 4.1, 5.0, 4.1, ACCENT, 2)
textbox(s, 0.9, 4.4, 11.5, 0.5, '[Name 1], [Name 2], [Name 3]', size=20, color=TEXT)
textbox(s, 0.9, 4.95, 11.5, 0.5, 'Politik, Klasse 10  |  Projekt Konfliktforschung  |  Oktober 2026', size=16,
        color=GREY)
set_notes(s, 1)

# 2 Lage
s = new_slide('Wo liegt der Gazastreifen?', 2)
map_image(s, 'gaza_overview', 0.6, 1.5, 6.3, 4.85)
caption(s, 0.6, 6.4, 6.3, 'Eigene Karte. Grenzen: OCHA / ArcGIS, Orte: OpenStreetMap')
bullets(s, 7.2, 1.55, 5.6, 5.2, [
    ('Lage: ', 'schmaler Küstenstreifen am Mittelmeer, Grenzen zu Israel und Ägypten'),
    ('Fläche: ', '365 km²'),
    ('Einwohner: ', 'ca. 2,1 Millionen (OCHA, Juli 2024), also rund 5.750 Menschen pro km²'),
    ('Größte Städte: ', 'Gaza-Stadt, Khan Yunis, Rafah'),
    ('Konfliktparteien: ', 'Israel (Armee: IDF) gegen Hamas und Palästinensischer Islamischer Dschihad'),
    ('Worum es geht: ', 'Sicherheit, Land und das Recht auf einen eigenen palästinensischen Staat'),
], size=18, space_after=12)
set_notes(s, 2)

# 3 Begriffe
s = new_slide('Wichtige Begriffe', 3)
bullets(s, 0.7, 1.55, 11.9, 5.2, [
    ('Asymmetrischer Konflikt: ', 'Die Gegner sind sehr unterschiedlich stark, z. B. eine reguläre Armee gegen '
                                  'bewaffnete Gruppen. Beide kämpfen mit unterschiedlichen Methoden.'),
    ('Humanitäres Völkerrecht: ', 'Regeln für den Krieg (z. B. Genfer Konventionen). Es schützt Zivilisten und gilt '
                                  'für alle Seiten, egal wer angefangen hat.'),
    ('UN-Sicherheitsrat und Veto: ', '15 Mitglieder, davon 5 ständige mit Vetorecht (USA, Russland, China, '
                                     'Frankreich, Großbritannien). Eine Resolution braucht 9 Stimmen und kein Veto '
                                     '(Art. 27 UN-Charta).'),
    ('Zwei-Staaten-Lösung: ', 'Israel und ein unabhängiger palästinensischer Staat existieren nebeneinander.'),
    ('Blockade: ', 'Israel und Ägypten kontrollieren seit 2007 streng, welche Waren und Personen nach Gaza hinein- '
                   'und herauskommen.'),
], size=20, space_after=20, line_spacing=1.05)
set_notes(s, 3)

# 4 Geschichte
s = new_slide('Wie ist der Konflikt entstanden?', 4)
history = [
    ('1948', 'Gründung Israels und Nakba', 'Krieg, Flucht und Vertreibung von Hunderttausenden Palästinensern'),
    ('1967', 'Sechstagekrieg', 'Israel besetzt den Gazastreifen, das Westjordanland und weitere Gebiete'),
    ('2005', 'Abzug aus Gaza', 'Israel zieht seine Soldaten ab und räumt die Siedlungen'),
    ('2007', 'Machtübernahme der Hamas', 'Die Hamas kontrolliert Gaza, Israel und Ägypten verschärfen die Blockade'),
    ('7. Oktober 2023', 'Terrorangriff der Hamas', 'Über 1.000 Tote, Geiselnahmen; Israel beginnt den Krieg in Gaza'),
]
for i, (year, head, body) in enumerate(history):
    y = 1.5 + i * 0.98
    textbox(s, 0.7, y, 2.6, 0.5, year, size=20, bold=True, color=ACCENT)
    textbox(s, 3.4, y, 9.2, 0.42, head, size=19, bold=True)
    textbox(s, 3.4, y + 0.42, 9.2, 0.5, body, size=16, color=GREY)
    if i < len(history) - 1:
        hline(s, 0.7, y + 0.93, 12.6, y + 0.93, RULE, 0.5)
set_notes(s, 4)

# 5 Interessen
s = new_slide('Wer will was? Die Interessen der Beteiligten', 5)
columns = [
    ('Israelische Regierung', [
        'Sicherheit für Israel und seine Bevölkerung',
        'Hamas militärisch ausschalten',
        'Rückkehr der Geiseln',
        'Eigene Macht erhalten',
    ]),
    ('Palästinensische Bevölkerung', [
        'Schutz vor Gewalt und Vertreibung',
        'Ende von Besatzung und Blockade',
        'Eigener Staat, Bewegungsfreiheit',
        'Wiederaufbau und politische Mitsprache',
    ]),
    ('Hamas-Führung', [
        'Eigene Herrschaft in Gaza sichern',
        'Bewaffneter Kampf gegen Israel',
        'Islamistische Ordnung durchsetzen',
        'Einfluss in der Region behalten',
    ]),
]
for i, (head, items) in enumerate(columns):
    x = 0.7 + i * 4.1
    rect(s, x, 1.5, 3.8, 0.65, LIGHT)
    textbox(s, x + 0.1, 1.5, 3.6, 0.65, head, size=19, bold=True, color=BLUE, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, x, 2.4, 3.8, 3.5, items, size=18, space_after=14)
textbox(s, 0.7, 6.1, 11.9, 0.7,
        'Die Palästinensische Autonomiebehörde (PA) regiert im Westjordanland und ist nicht die Hamas. '
        'Auch innerhalb Israels und unter den Palästinensern gibt es sehr unterschiedliche Meinungen.',
        size=14, italic=True, color=GREY)
set_notes(s, 5)

# 6 Recht und Machtmittel
s = new_slide('Rechtliche Lage und Machtmittel', 6)
textbox(s, 0.7, 1.5, 5.8, 0.5, 'Was sagt das Recht?', size=20, bold=True, color=BLUE)
bullets(s, 0.7, 2.1, 5.8, 4.5, [
    ('Art. 51 UN-Charta: ', 'Israel beruft sich auf Selbstverteidigung (im Besatzungskontext umstritten)'),
    ('Humanitäres Völkerrecht: ', 'gilt für alle; Zivilisten schützen, Geiselnahmen sind verboten'),
    ('IGH (Staaten): ', 'Verfahren Südafrika gegen Israel, vorläufige Maßnahmen 2024'),
    ('IStGH (Personen): ', 'Haftbefehle 2024 gegen Netanjahu, Gallant und Hamas-Kommandeur Deif'),
    ['sub', 'Haftbefehle und Eilmaßnahmen sind noch keine Urteile'],
], size=17, space_after=12)
hline(s, 6.75, 1.55, 6.75, 6.4, RULE, 0.75)
textbox(s, 7.0, 1.5, 5.8, 0.5, 'Welche Mittel haben die Parteien?', size=20, bold=True, color=BLUE)
bullets(s, 7.0, 2.1, 5.8, 4.5, [
    ('Israel: ', 'moderne Armee, Luftwaffe, Raketenabwehr; Waffen und Geld vor allem aus den USA'),
    ('Hamas / Islamischer Dschihad: ', 'Raketen, Tunnelsystem, Guerilla-Taktik, Geiseln als Druckmittel'),
    ('Iran: ', 'Geld, Waffen und Ausbildung für Hamas, Hisbollah und Huthi'),
    ('Katar: ', 'Kontakte zur Hamas, Hilfsgelder, Vermittler bei Verhandlungen'),
    ('Ägypten und USA: ', 'Vermittler, kontrollieren Grenzen bzw. Waffenlieferungen'),
], size=17, space_after=12)
set_notes(s, 6)

# 7 Verlauf und Zerstörung
s = new_slide('Verlauf des Krieges und Ausmaß der Zerstörung', 7)
bullets(s, 0.7, 1.45, 6.3, 2.3, [
    ('7. Oktober 2023: ', 'Angriff der Hamas, danach israelische Luftangriffe'),
    ('ab Ende Oktober 2023: ', 'Bodenoffensive der israelischen Armee'),
    ('2024/2025: ', 'Ausweitung auf Libanon (Hisbollah), Jemen (Huthi) und Iran'),
    ('seit Oktober 2025: ', 'Waffenruhe, die immer wieder verletzt wird'),
], size=15, space_after=6)
chart_data = CategoryChartData()
chart_data.categories = ['Zerstört', 'Schwer beschädigt', 'Mäßig beschädigt', 'Mögliche Schäden']
chart_data.add_series('Gebäude', (102067, 17421, 41895, 31429))
frame = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(3.55), Inches(6.5), Inches(2.9),
                           chart_data)
chart = frame.chart
chart.font.name = FONT
chart.font.size = Pt(12)
chart.has_legend = False
chart.has_title = True
chart.chart_title.text_frame.text = 'Beschädigte Gebäude im Gazastreifen (gesamt 192.812, ca. 78 %)'
title_run = chart.chart_title.text_frame.paragraphs[0].runs[0]
title_run.font.size = Pt(13)
title_run.font.bold = True
title_run.font.name = FONT
plot = chart.plots[0]
plot.gap_width = 60
plot.has_data_labels = True
plot.data_labels.number_format = '#,##0'
plot.data_labels.number_format_is_linked = False
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(12)
series = plot.series[0]
series.format.fill.solid()
series.format.fill.fore_color.rgb = ACCENT
chart.category_axis.reverse_order = True
chart.category_axis.tick_labels.font.size = Pt(12)
chart.category_axis.format.line.color.rgb = RULE
chart.value_axis.visible = False
chart.value_axis.has_major_gridlines = False
chart.value_axis.maximum_scale = 125000
caption(s, 0.7, 6.45, 6.4, 'Quelle: UNOSAT, Satellitenauswertung vom 8. Juli 2025 (vorläufig, nicht vor Ort geprüft)')
map_image(s, 'damage_governorates', 7.3, 1.5, 5.5, 4.4)
textbox(s, 7.3, 5.95, 5.5, 0.4, 'Betroffene Gebäude je Gouvernement', size=13, bold=True)
caption(s, 7.3, 6.3, 5.5, 'Absolute Zahlen, keine Prozentangaben. Quelle: UNOSAT, Stand 8.7.2025; eigene Karte')
set_notes(s, 7)

# 8 AKUF-Einordnung
s = new_slide('Welcher Kriegstyp ist das? Einordnung nach AKUF', 8)
textbox(s, 0.7, 1.5, 5.8, 0.5, 'Die fünf Kriegstypen nach AKUF', size=20, bold=True, color=BLUE)
bullets(s, 0.7, 2.1, 5.8, 4.2, [
    ('A  Antiregimekrieg: ', 'Kampf um die Regierung eines Landes'),
    ('B  Autonomie- / Sezessionskrieg: ', 'Kampf um Unabhängigkeit oder Selbstbestimmung'),
    ('C  Zwischenstaatlicher Krieg: ', 'Krieg zwischen zwei Staaten'),
    ('D  Dekolonisationskrieg: ', 'Befreiung von einer Kolonialmacht'),
    ('E  Sonstige Kriege', ''),
], size=16, space_after=12)
rect(s, 7.0, 1.5, 5.8, 4.9, LIGHT)
textbox(s, 7.25, 1.65, 5.4, 0.5, 'Unsere Einordnung: ein Mischtyp', size=20, bold=True, color=BLUE)
bullets(s, 7.25, 2.3, 5.4, 4.0, [
    ('Typ A: ', 'Es geht auch um das Ende oder den Erhalt der Hamas-Herrschaft in Gaza.'),
    ('Typ B: ', 'Die Palästinenser kämpfen um Selbstbestimmung und einen eigenen Staat.'),
    ('Typ C: ', 'Sobald Israel und der Iran sich direkt angreifen, wird es ein Krieg zwischen Staaten.'),
    'Die Hamas regiert in Gaza, ist aber kein Staat. Das macht die Einordnung schwierig.',
    'Besatzung ist bei AKUF kein eigener Typ.',
], size=16, space_after=10)
textbox(s, 0.7, 6.45, 12, 0.4, 'Das ist unsere eigene Einordnung, keine offizielle Bewertung durch die AKUF.',
        size=13, italic=True, color=GREY)
set_notes(s, 8)

# 9 UN-Sicherheitsrat (Tabelle)
s = new_slide('Was hat der UN-Sicherheitsrat gemacht?', 9)
rows = [
    ('Datum', 'Was passiert ist', 'Inhalt'),
    ('18.10.2023', 'Veto der USA', 'Entwurf Brasiliens für humanitäre Pausen scheitert'),
    ('25.10.2023', 'Vetos von Russland und China', 'Entwurf der USA scheitert (Streit um Selbstverteidigung und Waffenstillstand)'),
    ('25.03.2024', 'Resolution 2728', 'Sofortige Waffenruhe im Ramadan und Freilassung aller Geiseln; USA enthalten sich'),
    ('10.06.2024', 'Resolution 2735', 'Drei-Phasen-Plan: Waffenruhe und Austausch, dauerhaftes Ende der Kämpfe, Wiederaufbau'),
    ('17.11.2025', 'Resolution 2803', 'Unterstützt den Friedensplan, „Board of Peace“ als Übergangsverwaltung, internationale Stabilisierungstruppe'),
]
table_shape = s.shapes.add_table(len(rows), 3, Inches(0.7), Inches(1.5), Inches(11.9), Inches(4.0))
table = table_shape.table
table.columns[0].width = Inches(1.7)
table.columns[1].width = Inches(3.3)
table.columns[2].width = Inches(6.9)
for r, row in enumerate(rows):
    for c_idx, value in enumerate(row):
        cell = table.cell(r, c_idx)
        cell.text = value
        cell.margin_left = cell.margin_right = Inches(0.1)
        cell.margin_top = cell.margin_bottom = Inches(0.05)
        if r == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.name = FONT
                run.font.size = Pt(14)
                run.font.bold = (r == 0 or c_idx == 1)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else TEXT
                run._r.get_or_add_rPr().set('lang', 'de-DE')
bullets(s, 0.7, 5.75, 11.9, 1.1, [
    ('Vermittler: ', 'USA, Ägypten und Katar'),
    ('Hilfe vor Ort: ', 'UNRWA (Versorgung, Schulen) und OCHA (Koordination der humanitären Hilfe)'),
], size=15, space_after=4)
set_notes(s, 9)

# 10 Lösungswege
s = new_slide('Wie könnte eine Lösung aussehen?', 10)
bullets(s, 0.7, 1.55, 11.9, 5.3, [
    ('Anreize für Verhandlungen', ''),
    ['sub', 'Geisel- und Gefangenenaustausch war bisher der wichtigste Hebel'],
    ['sub', 'Humanitäre Not, militärische Kosten und internationaler Druck erhöhen den Einigungsdruck'],
    ('Vereinbarungen, die überprüfbar sind', ''),
    ['sub', 'Waffenruhe, die eingehalten wird; Schutz von Zivilisten; Hilfe ohne politische Bedingungen'],
    ['sub', 'Klare Regeln, was bei Verstößen passiert'],
    ('Politische Perspektive', ''),
    ['sub', 'Wiederaufbau, eine anerkannte palästinensische Vertretung und Sicherheit für Israel'],
    ['sub', 'Langfristig: zwei Staaten nebeneinander'],
    ('Zivilgesellschaft einbeziehen', ''),
    ['sub', 'Opferverbände, Friedensinitiativen und lokale Hilfsorganisationen schaffen Vertrauen'],
], size=21, sub_size=17, space_after=8, space_before=14)
set_notes(s, 10)

# 11 Bewertung
s = new_slide('Unsere Bewertung', 11)
textbox(s, 0.7, 1.5, 5.6, 0.5, 'Galtung: negativer und positiver Frieden', size=19, bold=True, color=BLUE)
bullets(s, 0.7, 2.1, 5.6, 2.8, [
    ('Negativer Frieden: ', 'Es wird nicht mehr geschossen. Die Waffenruhe seit Oktober 2025 wird aber immer '
                            'wieder gebrochen, also ist nicht einmal das erreicht.'),
    ('Positiver Frieden: ', 'Sicherheit, Rechte, Versorgung und Selbstbestimmung für alle; strukturelle Gewalt '
                            'wird abgebaut. Davon ist Gaza weit entfernt.'),
], size=16, space_after=10)
hline(s, 6.55, 1.55, 6.55, 4.85, RULE, 0.75)
textbox(s, 6.8, 1.5, 6.0, 0.5, 'Senghaas: sechs Bausteine für dauerhaften Frieden', size=19, bold=True, color=BLUE)
bullets(s, 6.8, 2.1, 6.0, 2.8, [
    'Staatliches Gewaltmonopol',
    'Rechtsstaatlichkeit',
    'Demokratische Beteiligung',
    'Soziale Gerechtigkeit',
    'Konstruktive Konfliktkultur',
    'Gegenseitige Abhängigkeit und Affektkontrolle',
], size=16, numbered=True, space_after=5)
rect(s, 0.7, 5.05, 11.9, 1.75, LIGHT)
textbox(s, 0.95, 5.15, 11.4, 0.4, 'Unser Fazit', size=18, bold=True, color=BLUE)
textbox(s, 0.95, 5.55, 11.4, 1.2,
        'Resolution 2803 ist ein möglicher Rahmen für den Übergang, aber noch kein Frieden. Von den sechs '
        'Bausteinen fehlen in Gaza vor allem eine legitime Kontrolle der Gewalt, Rechtsschutz, Versorgung und '
        'politische Beteiligung. Eine Lösung muss Israelis vor Angriffen schützen und Palästinensern gleiche '
        'Rechte und Selbstbestimmung geben.', size=15, line_spacing=1.05)
set_notes(s, 11)

# 12 Quellen
s = new_slide('Quellen', 12)
textbox(s, 0.7, 1.4, 11.9, 0.4, 'Alle Internetquellen abgerufen am 14.10.2026, Geodaten am 15.10.2026.',
        size=13, italic=True, color=GREY)
sources = [
    ('UN-Sicherheitsrat: Übersicht der Resolutionen', 'https://main.un.org/securitycouncil/en/content/resolutions-0'),
    ('UN-Sicherheitsrat: Resolution 2728 (2024)', 'https://docs.un.org/S/RES/2728(2024)'),
    ('UN-Sicherheitsrat: Resolution 2735 (2024)', 'https://docs.un.org/S/RES/2735(2024)'),
    ('UN-Sicherheitsrat: Resolution 2803 (2025)', 'https://digitallibrary.un.org/record/4093207'),
    ('UNOSAT: Gaza Strip Comprehensive Damage Assessment, 8.7.2025', 'https://unosat.org/products/4165'),
    ('AKUF (Universität Hamburg): Kriegsdefinition und Kriegstypologie',
     'https://www.wiso.uni-hamburg.de/fachbereich-sowi/ueber-den-fachbereich/fachgebiete/akuf/kriegsdefinition.html'),
    ('bpb: Der Gazastreifen im Nahostkonflikt (APuZ 2024)',
     'https://www.bpb.de/shop/zeitschriften/apuz/naher-osten-2024/546377/der-gazastreifen-im-nahostkonflikt/'),
    ('bpb: Geschichte des Nahostkonflikts',
     'https://www.bpb.de/themen/naher-mittlerer-osten/israel/45042/geschichte-des-nahostkonflikts/'),
    ('bpb: Frieden (Galtung, Senghaas)', 'https://www.bpb.de/lernen/politische-bildung/193093/frieden/?p=0'),
    ('SWP: Krieg in Nahost', 'https://www.swp-berlin.org/fokusthema/krieg-in-nahost'),
    ('OCHA: Gaza Strip access and movement, Juli 2024',
     'https://www.ochaopt.org/content/gaza-strip-access-and-movement-july-2024'),
    ('Geodaten für die Karten: ArcGIS GAZA_E (OCHA-Gouvernements, OpenStreetMap-Orte)',
     'https://services3.arcgis.com/tuNLpt6Wfhd22qmO/ArcGIS/rest/services/GAZA_E/FeatureServer'),
]
box = s.shapes.add_textbox(Inches(0.7), Inches(1.85), Inches(11.9), Inches(5.0))
tf = box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_right = Inches(0.05)
for i, (label, url) in enumerate(sources):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(8)
    _add_bullet(p, 0, numbered=True)
    r = p.add_run()
    r.text = label + ': '
    _style_run(r, 13, False, False, TEXT)
    r2 = p.add_run()
    r2.text = url
    _style_run(r2, 12, False, False, GREY)
    r2.hyperlink.address = url
textbox(s, 0.7, 6.55, 11.9, 0.4,
        'Weitere Quellen (Gerichtsverfahren, Vetos, Hilfsorganisationen) stehen in den Notizen zu den Folien.',
        size=12, italic=True, color=GREY)
set_notes(s, 12)


# ---------------------------------------------------------------- Kontrolle und Speichern

assert len(prs.slides) == TOTAL
for i, slide in enumerate(prs.slides, 1):
    assert slide.notes_slide.notes_text_frame.text.strip()
    assert f'Sprecher: [Name {(i - 1) // 4 + 1}]' in slide.notes_slide.notes_text_frame.text
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, (i, shape.name, 'negative')
        assert shape.left + shape.width <= prs.slide_width + 2, (i, shape.name, 'rechts')
        assert shape.top + shape.height <= prs.slide_height + 2, (i, shape.name, 'unten')
        if shape.has_text_frame:
            assert shape.text_frame.word_wrap

parser = argparse.ArgumentParser()
parser.add_argument('--draft', action='store_true', help='nach build/ statt output/ schreiben')
args = parser.parse_args()
target = BUILD / 'candidate_redesign.pptx' if args.draft else OUT / 'akuf_gazastreifen_neugestaltet.pptx'
prs.save(target)
(BUILD / 'image_checks.json').write_text(json.dumps(IMAGE_CHECKS, indent=2))
print(f'{TOTAL} Folien gespeichert: {target}')

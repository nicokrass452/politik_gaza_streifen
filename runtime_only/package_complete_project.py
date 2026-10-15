"""Assemble and verify the complete, portable Gaza presentation archive."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from hashlib import sha256
from collections import Counter
from datetime import datetime, timezone
import ast,csv,io,json,shutil,sys,importlib.metadata
from pptx import Presentation

WORK=Path(__file__).resolve().parent
SRC=WORK/'gaza'
ROOT=WORK/'complete_archive'
OUT=WORK/'deliverables'
ROOT.mkdir(exist_ok=True);OUT.mkdir(exist_ok=True)


def copy(source,relative):
    source=Path(source);target=ROOT/relative
    target.parent.mkdir(parents=True,exist_ok=True)
    shutil.copyfile(source,target)
    return relative

def put(relative,text):
    target=ROOT/relative;target.parent.mkdir(parents=True,exist_ok=True);target.write_text(text,encoding='utf-8')

def put_json(relative,data):put(relative,json.dumps(data,ensure_ascii=False,indent=2)+'\n')
def put_csv(relative,fields,rows):
    out=io.StringIO(newline='');w=csv.DictWriter(out,fieldnames=fields);w.writeheader();w.writerows(rows)
    put(relative,out.getvalue())

# Two finished presentations, both retained under their delivered filenames.
copy(SRC/'redesign/output/akuf_gazastreifen_neugestaltet.pptx','presentations/akuf_gazastreifen_neugestaltet.pptx')
copy(SRC/'delivery/akuf_gazastreifen_praesentation.pptx','presentations/akuf_gazastreifen_praesentation_original.pptx')

# Original and revised authoring projects preserve their relative asset paths.
copy(SRC/'build_presentation.py','projects/original/build_presentation.py')
copy(SRC/'requirements.txt','projects/original/requirements.txt')
for p in sorted((SRC/'assets').iterdir()):
    if p.is_file() and not p.name.startswith('.'):
        copy(p,'projects/original/assets/'+p.name)
for name in ['build_redesign.py','generate_graphics.py','rasterize_svgs.py','requirements.txt','README.md']:
    copy(SRC/'redesign'/name,'projects/redesign/'+name)
for p in sorted((SRC/'redesign/assets').rglob('*')):
    if p.is_file() and not any(x.startswith('.') for x in p.relative_to(SRC/'redesign/assets').parts):
        copy(p,'projects/redesign/assets/'+p.relative_to(SRC/'redesign/assets').as_posix())

# Preserve all task-authored JavaScript, including optional environment-specific checks.
copy(SRC/'build/finalize.mjs','runtime_only/original_finalize.mjs')
for name in ['finalize.mjs','render_svgs.mjs']:
    copy(SRC/'redesign'/name,'runtime_only/redesign_'+name)
put('runtime_only/README.md','''# Runtime-specific authoring utilities\n\nThese are the exact task-authored JavaScript utilities used during authoring. They reference the original ChatGPT presentation runtime and original workspace paths. They are included for completeness, but are optional and are not portable standalone programs.\n\nUse the Python workflows under `projects/` to regenerate the presentation on your own machine. The third-party presentation runtime and authentication/upload utilities are not bundled.\n''')

copy(Path(__file__), 'runtime_only/package_complete_project.py')

# Validation records and all final slide previews.
for name in ['validation.json','validation-final.json','image_checks.json','notes_manifest.json']:
    copy(SRC/'build'/name,'validation/original/'+name)
for name in ['final_validation.json','geo_validation.json','image_checks.json']:
    copy(SRC/'redesign/build'/name,'validation/redesign/'+name)
for version,directory in [('original',SRC/'build/rendered-final'),('redesign',SRC/'redesign/build/reviewed')]:
    for p in sorted(directory.glob('slide-*.png')):copy(p,f'previews/{version}/'+p.name)

# Include the user's reusable layout draft as a reference, without installing it.
for p in sorted((WORK/'skills/ppt-basic-layout-creator').rglob('*')):
    if p.is_file():copy(p,'layout_skill_draft/ppt-basic-layout-creator/'+p.relative_to(WORK/'skills/ppt-basic-layout-creator').as_posix())
put('layout_skill_draft/STATUS.md','''# Layout skill status\n\nThis is the locally validated layout skill draft, including proportional image fitting. It was not installed or pushed to a Skills Git repository. It is included here as a reusable project reference.\n''')

# Extract the complete original citation registry without executing the builder.
source_dict={}
for node in ast.parse((SRC/'build_presentation.py').read_text()).body:
    if isinstance(node,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='SOURCES' for t in node.targets):
        source_dict=ast.literal_eval(node.value)
notes_manifest=json.loads((SRC/'build/notes_manifest.json').read_text())
source_rows=[]
raw={
    'map':['projects/original/assets/ocha_gaza_2024.pdf'],
    'unosat':['projects/original/assets/unosat_2025-07-08.pdf'],
    'unosatpdf':['projects/original/assets/unosat_2025-07-08.pdf'],
}
for key,(title,url) in source_dict.items():
    files=raw.get(key,[])
    if key in ['r2728','r2735','r2803']:
        no={'r2728':'2728_2024','r2735':'2735_2024','r2803':'2803_2025'}[key]
        p=ROOT/f'research/primary_documents/resolution_{no}_en.pdf'
        if p.exists():files=files+[p.relative_to(ROOT).as_posix()]
    source_rows.append(dict(id=key,title=title,url=url,research_access_date='2026-10-14',used_in_slides=','.join(str(n['slide']) for n in notes_manifest if key in n['sources']),archive_representation='stored_document_and_notes' if files else 'citation_and_research_notes',local_files=';'.join(files)))
new_sources=[
 ('geo_governorates','Gaza governorate boundaries, GAZA_E layer 34','https://services3.arcgis.com/tuNLpt6Wfhd22qmO/ArcGIS/rest/services/GAZA_E/FeatureServer/34','2,7','projects/redesign/assets/gaza_governorates.geojson'),
 ('geo_places','OpenStreetMap places, GAZA_E layer 32','https://services3.arcgis.com/tuNLpt6Wfhd22qmO/ArcGIS/rest/services/GAZA_E/FeatureServer/32','2','projects/redesign/assets/gaza_places.geojson'),
 ('geo_outline','Gaza Strip boundary, GAZA_E layer 35','https://services3.arcgis.com/tuNLpt6Wfhd22qmO/ArcGIS/rest/services/GAZA_E/FeatureServer/35','','projects/redesign/assets/gaza_ocha_adm1.geojson'),
 ('geo_metadata','GAZA_E service metadata','https://services3.arcgis.com/tuNLpt6Wfhd22qmO/ArcGIS/rest/services/GAZA_E/FeatureServer?f=pjson','','projects/redesign/assets/arcgis_metadata.json'),
 ('geoboundaries','geoBoundaries PSE ADM1, alternative boundary data','https://www.geoboundaries.org/api/current/gbOpen/PSE/ADM1/','','projects/redesign/assets/pse_adm1.geojson'),
 ('lucide','Lucide icon sources and license','https://lucide.dev/license','1–12','projects/redesign/assets/icons/LICENSE'),
 ('fluter_map','fluter/bpb Gaza map, considered but not used in final deck','https://www.fluter.de/gazastreifen-geschichte','','projects/redesign/assets/map_fluter.jpg'),
 ('ocha_snapshot','OCHA Reported Impact Snapshot, 16 July 2025','https://www.ochaopt.org/sites/default/files/Gaza_Reported_Impact_Snapshot_16_July_2025.pdf','','projects/original/assets/ocha_snapshot_2025-07-16.pdf'),
]
for key,title,url,slides,files in new_sources:
    source_rows.append(dict(id=key,title=title,url=url,research_access_date='2026-10-15' if key!='ocha_snapshot' else '2026-10-14',used_in_slides=slides,archive_representation='stored_document_or_data',local_files=files))
put_json('research/sources.json',source_rows)
put_csv('research/sources.csv',list(source_rows[0]),source_rows)
copy(SRC/'quellen.md','research/original_bibliography_de.md')
index=['# Research sources','', 'Original subject research: 14 October 2026. Geographic/design research: 15 October 2026. This registry includes both used and considered sources. `citation_and_research_notes` means the full web article is not stored offline; the source URL and the prepared research notes are included.','']
for r in source_rows:
    index.extend([f"## {r['id']}: {r['title']}",r['url'],f"Research access date: {r['research_access_date']}",f"Slides: {r['used_in_slides'] or 'not used in the final slides'}",f"Archive representation: {r['archive_representation']}"])
    if r['local_files']:index.append('Local files: '+r['local_files'])
    index.append('')
put('research/SOURCES.md','\n'.join(index))

# Full slide text and prepared speaking notes are reviewable outside PowerPoint.
all_notes={}
for label,path in [('original',ROOT/'presentations/akuf_gazastreifen_praesentation_original.pptx'),('redesign',ROOT/'presentations/akuf_gazastreifen_neugestaltet.pptx')]:
    prs=Presentation(path);rows=[];md=[f'# Sprechernotizen / {label}','']
    for i,s in enumerate(prs.slides,1):
        text='\n'.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())
        note=s.notes_slide.notes_text_frame.text
        rows.append(dict(slide=i,speaker=f'[Name {(i-1)//4+1}]',slide_text=text,speaker_notes=note))
        md += [f'## Folie {i:02d} / [Name {(i-1)//4+1}]','', '### Folientext','',text,'','### Sprechernotizen und Quellen','',note,'']
    all_notes[label]=rows;put_json(f'research/notes/{label}.json',rows);put(f'research/notes/{label}.md','\n'.join(md))
put_csv('research/slide_assignment.csv',['slide','speaker','target_seconds','original_source_ids'],[dict(slide=n['slide'],speaker=f"[Name {n['speaker']}]",target_seconds=n['seconds'],original_source_ids=','.join(n['sources'])) for n in notes_manifest])

# Machine-readable copies of all numerical inputs used in the SVG graphics.
damage=json.loads((SRC/'redesign/assets/damage_data.json').read_text())
names={'PS0255':'North Gaza','PS0260':'Gaza','PS0265':'Deir Al-Balah','PS0270':'Khan Younis','PS0275':'Rafah'}
put_csv('research/data/damage_by_governorate_2025-07-08.csv',['pcode','governorate','affected_structures','imagery_date','source_id'],[dict(pcode=k,governorate=names[k],affected_structures=v,imagery_date='2025-07-08',source_id='unosat') for k,v in damage['counts_by_pcode'].items()])
put_csv('research/data/damage_categories_2025-07-08.csv',['category_de','structure_count','imagery_date','source_id'],[dict(category_de=k,structure_count=v,imagery_date='2025-07-08',source_id='unosat') for k,v in damage['categories']])
put_json('research/data/damage_statistics.json',damage)
put_csv('research/data/overview_population.csv',['metric','value','unit','data_period','qualification','source_id'],[
 dict(metric='population_estimate',value=2100000,unit='people',data_period='2024-07',qualification='rounded OCHA estimate, not a current population distribution',source_id='map'),
 dict(metric='area',value=365,unit='km2',data_period='2024-07',qualification='rounded geographic area',source_id='map'),
 dict(metric='average_density_displayed',value=5750,unit='people_per_km2',data_period='2024-07',qualification='rounded population/area calculation, not local density',source_id='map')])
put_csv('research/data/unsc_milestones.csv',['resolution','adoption_date','votes_for','votes_against','abstentions','source_id'],[
 dict(resolution='2728 (2024)',adoption_date='2024-03-25',votes_for=14,votes_against=0,abstentions=1,source_id='r2728'),
 dict(resolution='2735 (2024)',adoption_date='2024-06-10',votes_for=14,votes_against=0,abstentions=1,source_id='r2735'),
 dict(resolution='2803 (2025)',adoption_date='2025-11-17',votes_for=13,votes_against=0,abstentions=2,source_id='r2803')])
put_csv('research/data/akuf_typology.csv',['type','label_de','source_id'],[dict(type=a,label_de=b,source_id='akuf') for a,b in [('A','Antiregimekrieg'),('B','Autonomie- und Sezessionskrieg'),('C','Zwischenstaatlicher Krieg'),('D','Dekolonisationskrieg'),('E','Sonstige Kriege')]])
put('research/DATA_NOTES.md','''# Data notes\n\n- UNOSAT building data are a historical snapshot based on imagery dated 8 July 2025, version 1, published 6 August 2025. The count 192,812 includes destroyed, severely damaged, moderately damaged and possibly damaged structures. The reported 78% is rounded.\n- Governorate counts are absolute counts, not within-governorate damage percentages. They must not be read as current control zones or individual damage locations. The supplied dataset contains five governorate polygons and city points, not the full UNOSAT building-level damage dataset.\n- The population estimate is historical (OCHA, July 2024). The displayed density is a rounded arithmetic average, not a population-density raster.\n- Boundary areas are checked in EPSG:32636 (UTM 36N). The saved geometry area is about 364.88 km², consistent with the rounded 365 km². The geographic map is for orientation and does not show current access routes or control zones.\n- The presentation's mixed AKUF classification and its final judgement are explicitly the group's analytical interpretation, not an official AKUF case code.\n- The archive preserves the prepared presentation and research snapshots; it does not silently update the facts to a later date.\n- Every original research source has a registry entry. Full PDFs and geographic files are included where available; other source pages are represented by citations and the prepared notes, not full offline article copies.\n''')

# Portable helper for fetching the same public raw source files again, if needed.
fetch_records=[
 dict(path='original/ocha_gaza_2024.pdf',url='https://www.ochaopt.org/sites/default/files/Gaza_A0_2024_july-2024.pdf',kind='pdf'),
 dict(path='original/unosat_2025-07-08.pdf',url=source_dict['unosatpdf'][1],kind='pdf'),
 dict(path='original/ocha_snapshot_2025-07-16.pdf',url='https://www.ochaopt.org/sites/default/files/Gaza_Reported_Impact_Snapshot_16_July_2025.pdf',kind='pdf'),
 dict(path='geodata/gaza_governorates.geojson',url='https://services3.arcgis.com/tuNLpt6Wfhd22qmO/ArcGIS/rest/services/GAZA_E/FeatureServer/34/query',params={'where':'1=1','outFields':'*','outSR':4326,'f':'geojson'},kind='geojson'),
 dict(path='geodata/gaza_places.geojson',url='https://services3.arcgis.com/tuNLpt6Wfhd22qmO/ArcGIS/rest/services/GAZA_E/FeatureServer/32/query',params={'where':'1=1','outFields':'*','outSR':4326,'f':'geojson'},kind='geojson'),
]
fetch_records.extend([dict(path='primary_documents/resolution_'+no+'_'+year+'_en.pdf',url='https://digitallibrary.un.org/record/'+record+'/files/S_RES_'+no+'_%28'+year+'%29-EN.pdf',kind='pdf') for no,year,record in [('2728','2024','4042189'),('2735','2024','4051310'),('2803','2025','4093207')]])
put_json('research/public_download_manifest.json',fetch_records)
put('research/fetch_public_data.py', '''"""Optionally re-download public source files. Requires requests. Never overwrites archived evidence."""\nfrom pathlib import Path\nfrom hashlib import sha256\nimport json, requests\nroot=Path(__file__).resolve().parent\nrecords=json.loads((root/'public_download_manifest.json').read_text())\nout=root/'fresh_downloads';out.mkdir(exist_ok=True)\nlog=[]\nfor item in records:\n    try:\n        r=requests.get(item['url'],params=item.get('params'),timeout=45);r.raise_for_status()\n        if item['kind']=='pdf' and not r.content.startswith(b'%PDF'):raise ValueError('Response is not a PDF')\n        if item['kind']=='geojson' and r.json().get('type')!='FeatureCollection':raise ValueError('Response is not a GeoJSON collection')\n        p=out/item['path'];p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(r.content)\n        log.append(dict(path=item['path'],status='saved',sha256=sha256(r.content).hexdigest()))\n    except Exception as e:log.append(dict(path=item['path'],status='failed',error=str(e)))\n(out/'download_log.json').write_text(json.dumps(log,indent=2));print(json.dumps(log,indent=2))\n''')

# Run instructions and scope are explicit so the archive stays useful off this machine.
put('README.md','''# Complete Gaza / AKUF presentation project\n\nThis archive contains the original and redesigned German school presentations, all task-authored project code, all saved research assets and geographic data, the SVG graphics, full speaker notes, final slide previews and validation records.\n\n## Start here\n\n- `presentations/akuf_gazastreifen_neugestaltet.pptx`: latest cream, charcoal and terracotta design; use this for the presentation.\n- `presentations/akuf_gazastreifen_praesentation_original.pptx`: the earlier slate-and-blue version, retained for comparison.\n- Both finished decks have 12 slides, full German notes and four slides per speaker. Replace `[Name 1]`, `[Name 2]` and `[Name 3]` in PowerPoint before presenting.\n\n## Contents\n\n| Folder | Contents |\n|---|---|\n| `projects/original/` | Original Python builder, requirements, original OCHA and UNOSAT PDFs and PNGs |\n| `projects/redesign/` | GeoPandas/SVG generator, PowerPoint builder, portable SVG rasterizer, all SVGs, previews, geographic inputs, metadata, icons and icon license |\n| `research/` | Full source registry in Markdown/CSV/JSON, full slide text and speaking notes for both versions, statistical CSV/JSON inputs and source-document download records |\n| `previews/` | All 12 final slide previews for each finished version |\n| `validation/` | Saved geometry, image-fit, slide-count and package validation reports |\n| `runtime_only/` | Task-authored JavaScript used with the original presentation runtime |\n| `layout_skill_draft/` | The original user-defined layout skill draft, including image fitting |\n\nRaw map PDFs are stored with `projects/original/assets/`; geographic files and graphic metadata are under `projects/redesign/assets/`. `research/SOURCES.md` maps sources to their stored files or prepared notes. The archive is a complete saved project, not an offline mirror of every website visited. Some web sources are represented by URLs and notes rather than complete article copies.\n\n## Regenerate the redesigned deck\n\nFrom `projects/redesign/`:\n\n```bash\npip install -r requirements.txt\npython generate_graphics.py\npython rasterize_svgs.py\npython build_redesign.py\n```\n\nThe generated PowerPoint appears in that project's `output/` folder. All inputs are already included, so no data download is needed for this workflow. To rebuild using the included, already-rendered SVG/PNG assets, only the PowerPoint builder is needed.\n\nNimbus Roman and Nimbus Sans were used for the redesigned deck. They are not bundled. PowerPoint may substitute fonts on systems where they are unavailable; inspect layout after substitution. CairoSVG may also require a platform-specific Cairo installation.\n\n## Regenerate the earlier version\n\nFrom `projects/original/`:\n\n```bash\npip install -r requirements.txt\npython build_presentation.py\n```\n\nThe map PDFs and image files are included. The original builder's download paths are only used when required map assets are missing.\n\n## Research dates and data scope\n\n- Subject research: 14 October 2026.\n- Geographic and design research: 15 October 2026.\n- UNOSAT damage imagery: 8 July 2025, version 1.\n- Population estimate: OCHA, July 2024.\n\nThe historical figures retain their original dates. `research/DATA_NOTES.md` explains units, rounding and interpretation. Raw governorate and city geometry, source metadata, numerical inputs, prepared research notes and citations are included. The full building-level UNOSAT damage point dataset was not used and is not part of this archive.\n\nThe small icons retain their Lucide license. The considered fluter map and alternative boundary files are included as research assets, although they were not used in the final redesigned slides. Bibliographic sources remain subject to their original terms.\n\nThe optional runtime-specific JavaScript requires the original authoring runtime and original path setup. Generic installed packages, package caches, credentials, authentication helpers, duplicate temporary exports and nested copies of previous ZIP bundles are omitted. Project dependencies are listed in the requirements files.\n\n## Integrity\n\n`FILE_INDEX.csv` lists every payload file's size and SHA-256 checksum. `SHA256SUMS.txt` provides a checksum list. `ARCHIVE_SUMMARY.json` records file counts and archive coverage.\n''')
put('START_HERE_zh.md','''# 完整项目包\n\n最新演示文稿：`presentations/akuf_gazastreifen_neugestaltet.pptx`。\n\n本包包含两版正式 PPT、全部编写的项目代码、已保存的研究资料、原始 OCHA／UNOSAT PDF、GeoJSON 地理数据、SVG 图形、统计 CSV／JSON、两版德语讲稿、24 张最终幻灯片预览和校验记录。\n\n完整来源索引见 `research/SOURCES.md`。部分网页未保存全文，索引保留了网址和研究讲稿。原始地图文件在 `projects/original/assets/`，重设计数据和图形在 `projects/redesign/assets/`。\n\n复现步骤见根目录 `README.md`。通用依赖通过 requirements 安装，不需要把本机依赖目录复制到其他电脑。\n''')

# Capture relevant installed versions as a reference, without shipping libraries or caches.
versions={}
for name in ['python-pptx','Pillow','requests','geopandas','shapely','pyproj','pyogrio','pandas','lxml']:
    try:versions[name]=importlib.metadata.version(name)
    except importlib.metadata.PackageNotFoundError:versions[name]='not found in active interpreter'
put_json('validation/environment_reference.json',dict(python=sys.version,packages=versions,portable_entrypoints=['projects/original/build_presentation.py','projects/redesign/build_redesign.py','projects/redesign/generate_graphics.py']))

# A complete inventory is included in the ZIP; validation doesn't alter the decks.
files=sorted(p for p in ROOT.rglob('*') if p.is_file() and p.name not in ['FILE_INDEX.csv','SHA256SUMS.txt','ARCHIVE_SUMMARY.json'])
summary=dict(created_date='2026-10-15',finished_presentations=2,slide_previews=24,source_registry_entries=len(source_rows),saved_research_snapshot_dates=['2026-10-14','2026-10-15'],svg_asset_count=len(list((ROOT/'projects/redesign/assets').rglob('*.svg'))),geojson_count=len(list(ROOT.rglob('*.geojson'))),payload_file_count=len(files),payload_bytes=sum(p.stat().st_size for p in files),notes_versions=['original','redesign'])
put_json('ARCHIVE_SUMMARY.json',summary)
files=sorted(p for p in ROOT.rglob('*') if p.is_file() and p.name not in ['FILE_INDEX.csv','SHA256SUMS.txt'])
rows=[dict(path=p.relative_to(ROOT).as_posix(),bytes=p.stat().st_size,sha256=sha256(p.read_bytes()).hexdigest()) for p in files]
put_csv('FILE_INDEX.csv',['path','bytes','sha256'],rows)
put('SHA256SUMS.txt','\n'.join(f"{r['sha256']}  {r['path']}" for r in rows)+'\n')
archive=OUT/'akuf_gazastreifen_complete_project.zip'
with ZipFile(archive,'w',ZIP_DEFLATED,compresslevel=9) as z:
    for p in sorted(ROOT.rglob('*')):
        if p.is_file():z.write(p,'akuf_gazastreifen_complete/'+p.relative_to(ROOT).as_posix())
with ZipFile(archive) as z:
    bad=z.testzip();assert bad is None,bad
    for r in rows:
        data=z.read('akuf_gazastreifen_complete/'+r['path'])
        assert sha256(data).hexdigest()==r['sha256'],r['path']
    assert any(n.endswith('presentations/akuf_gazastreifen_neugestaltet.pptx') for n in z.namelist())
    count=len(z.namelist())
print(json.dumps(dict(archive=str(archive),size_bytes=archive.stat().st_size,size_mib=round(archive.stat().st_size/1048576,2),files=count,crc_and_checksum_validation='passed',sources=len(source_rows)),indent=2))
